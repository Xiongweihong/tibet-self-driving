import openpyxl
import json

# Read 49 regions from 西藏自驾.xlsx
wb = openpyxl.load_workbook('西藏自驾.xlsx')
ws = wb.active

regions_list = []
for r in range(3, ws.max_row + 1):
    num = ws.cell(r, 1).value
    pref = ws.cell(r, 2).value
    cty = ws.cell(r, 3).value
    spots = ws.cell(r, 4).value
    food = ws.cell(r, 5).value
    hotel = ws.cell(r, 6).value
    weather = ws.cell(r, 7).value
    pet = ws.cell(r, 8).value

    if pref and cty:
        # Determine category tag
        stage_tag = "川西翡翠阶梯"
        if "雅安" in pref or "乐山" in pref:
            stage_tag = "川西雅安·大渡河"
        elif "甘孜" in pref:
            if cty in ["德格县", "甘孜县", "炉霍县", "道孚县", "丹巴县"]:
                stage_tag = "川藏北线G317"
            else:
                stage_tag = "川西甘孜G318"
        elif "昌都" in pref:
            if cty in ["丁青县", "类乌齐县", "卡若区", "江达县"]:
                stage_tag = "川藏北线G317"
            else:
                stage_tag = "藏东昌都G318"
        elif "林芝" in pref:
            stage_tag = "林芝雪域江南"
        elif "山南" in pref:
            stage_tag = "藏南雅江摇篮"
        elif "拉萨" in pref:
            stage_tag = "拉萨圣城及周边"
        elif "阿里" in pref:
            stage_tag = "西极阿里全境"
        elif "那曲" in pref:
            stage_tag = "藏北内陆湖泊"
        elif "日喀则" in pref:
            stage_tag = "远眺合规区"

        regions_list.append({
            "id": int(num) if num else len(regions_list) + 1,
            "pref": str(pref),
            "cty": str(cty),
            "stage": stage_tag,
            "spots": str(spots or ""),
            "food": str(food or ""),
            "hotel": str(hotel or ""),
            "weather": str(weather or ""),
            "pet": str(pet or "")
        })

print(f"Extracted {len(regions_list)} regions from Excel.")

# Export to json file for index.html injection
with open('regions_data.json', 'w', encoding='utf-8') as f:
    json.dump(regions_list, f, ensure_ascii=False, indent=2)

print("Saved regions_data.json successfully.")
